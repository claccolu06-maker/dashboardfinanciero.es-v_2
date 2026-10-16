"""Lectores de archivos: CSV/TXT, Excel (.xlsx) y PDF -> texto para el motor."""
from __future__ import annotations

import datetime
import io
import zipfile
from typing import List, Tuple

from .limites import (
    ImportacionError, MAX_BYTES, MAX_LINEAS, MAX_PAGINAS_PDF, MAX_RATIO_COMPRESION,
    MAX_XLSX_DESCOMPRIMIDO, MAX_XLSX_ENTRADAS,
)

_FIRMA_PDF = b"%PDF-"
_FIRMA_ZIP = b"PK\x03\x04"
_FIRMA_OLE = b"\xd0\xcf\x11\xe0"       # .xls antiguo
MARCA_PAGINA = "@@PAGINA@@"            # separa páginas para quitar cabeceras/pies repetidos


def detectar_tipo(nombre: str, contenido: bytes) -> str:
    """'pdf' | 'xlsx' | 'texto' mirando primero la firma y luego la extensión."""
    cabecera = contenido[:8]
    if cabecera.startswith(_FIRMA_PDF):
        return "pdf"
    if cabecera.startswith(_FIRMA_ZIP):
        return "xlsx"
    if cabecera.startswith(_FIRMA_OLE):
        raise ImportacionError(
            "Los .xls antiguos no se pueden leer con seguridad. Guárdalo como .xlsx o .csv e inténtalo de nuevo.")
    ext = (nombre or "").lower().rsplit(".", 1)[-1] if "." in (nombre or "") else ""
    if ext in {"pdf"}:
        raise ImportacionError("El archivo no parece un PDF válido.")
    if ext in {"xlsx", "xlsm"}:
        raise ImportacionError("El archivo no parece un Excel (.xlsx) válido.")
    return "texto"


def decodificar_texto(contenido: bytes) -> str:
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return contenido.decode("utf-8", errors="replace")  # pragma: no cover


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _validar_zip_xlsx(contenido: bytes) -> None:
    """Rechaza zip bombs antes de que openpyxl descomprima nada."""
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            infos = z.infolist()
            if len(infos) > MAX_XLSX_ENTRADAS:
                raise ImportacionError("El Excel contiene demasiados elementos internos.")
            total = 0
            for i in infos:
                total += i.file_size
                if total > MAX_XLSX_DESCOMPRIMIDO:
                    raise ImportacionError("El Excel descomprimido es demasiado grande.")
                if i.file_size > 5 * 1024 * 1024 and i.compress_size and \
                        i.file_size / i.compress_size > MAX_RATIO_COMPRESION:
                    raise ImportacionError("El Excel tiene una relación de compresión sospechosa.")
            nombres = {i.filename for i in infos}
            if "xl/workbook.xml" not in nombres and "[Content_Types].xml" not in nombres:
                raise ImportacionError("El archivo no parece un Excel (.xlsx) válido.")
    except zipfile.BadZipFile:
        raise ImportacionError("El archivo no parece un Excel (.xlsx) válido.") from None


def _celda_a_texto(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return f"{v:.2f}"
    return " ".join(str(v).replace("\t", " ").replace("\n", " ").replace("\r", " ").split())


def texto_desde_xlsx(contenido: bytes) -> str:
    _validar_zip_xlsx(contenido)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportacionError("Falta la librería openpyxl en el servidor.") from exc
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ImportacionError("No se ha podido abrir el Excel. ¿Está protegido con contraseña o dañado?") from None
    mejor: List[str] = []
    try:
        for hoja in libro.worksheets:
            lineas: List[str] = []
            for fila in hoja.iter_rows(values_only=True):
                celdas = [_celda_a_texto(c) for c in fila]
                if any(celdas):
                    lineas.append("\t".join(celdas).rstrip("\t"))
                if len(lineas) > MAX_LINEAS:
                    raise ImportacionError(f"El Excel supera el máximo de {MAX_LINEAS:,} filas.".replace(",", "."))
            if len(lineas) > len(mejor):
                mejor = lineas
    finally:
        libro.close()
    if not mejor:
        raise ImportacionError("El Excel está vacío.")
    return "\n".join(mejor)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def textos_desde_pdf(contenido: bytes) -> List[str]:
    """Devuelve una o dos extracciones del PDF (texto normal y modo 'layout')."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ImportacionError("Falta la librería pypdf en el servidor.") from exc
    try:
        lector = PdfReader(io.BytesIO(contenido))
        if lector.is_encrypted:
            try:
                if not lector.decrypt(""):
                    raise ImportacionError("El PDF está protegido con contraseña.")
            except ImportacionError:
                raise
            except Exception:
                raise ImportacionError("El PDF está protegido con contraseña.") from None
        paginas = lector.pages
        if len(paginas) > MAX_PAGINAS_PDF:
            raise ImportacionError(f"El PDF tiene {len(paginas)} páginas; el máximo es {MAX_PAGINAS_PDF}.")
        variantes: List[str] = []
        sep = f"\n{MARCA_PAGINA}\n"
        normal = sep.join((p.extract_text() or "") for p in paginas)
        if normal.replace(MARCA_PAGINA, "").strip():
            variantes.append(normal)
        try:
            layout = sep.join((p.extract_text(extraction_mode="layout") or "") for p in paginas)
            if layout.replace(MARCA_PAGINA, "").strip():
                variantes.append(layout)
        except TypeError:           # pypdf antiguo sin modo layout
            pass
    except ImportacionError:
        raise
    except Exception:
        raise ImportacionError("No se ha podido leer el PDF. ¿Está dañado?") from None
    if not variantes:
        raise ImportacionError(
            "El PDF no contiene texto seleccionable (¿es una imagen escaneada?). "
            "Descarga el extracto en CSV/Excel o copia y pega el texto desde la web del banco.")
    return variantes


def leer_archivo(nombre: str, contenido: bytes) -> Tuple[str, List[str]]:
    """Devuelve ``(tipo, [textos])`` listo para el motor."""
    if not contenido:
        raise ImportacionError("El archivo está vacío.")
    if len(contenido) > MAX_BYTES:
        raise ImportacionError(f"El archivo supera el máximo de {MAX_BYTES // (1024 * 1024)} MB.")
    tipo = detectar_tipo(nombre, contenido)
    if tipo == "pdf":
        return tipo, textos_desde_pdf(contenido)
    if tipo == "xlsx":
        return tipo, [texto_desde_xlsx(contenido)]
    texto = decodificar_texto(contenido)
    if "\x00" in texto[:2000]:
        raise ImportacionError("El archivo parece binario; sube un CSV, un Excel (.xlsx) o un PDF.")
    return tipo, [texto]
